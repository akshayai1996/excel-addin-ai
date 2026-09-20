"""Generate vigorous-test workbooks for the Excel AI add-in.

Outputs (in this folder):
  sales_test.xlsx — 600-row Sales table + formulas + Total row, Config sheet
                    with TaxRate named range, Raw pivot source, hidden Archive.
  messy_test.xlsx — dirty Data sheet (N/A, case mixes, blanks, duplicates,
                    unsorted) + empty Notes sheet.
Covers test plan phases A–E (Section 87 context).
"""
import random
from pathlib import Path

import pyopenxlsx
from pyopenxlsx import Font, Fill, Alignment

OUT = Path(__file__).resolve().parent
random.seed(42)

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
REGIONS = ["North", "South", "East", "West"]
PRODUCTS = ["Valves", "Pumps", "Pipes", "Flanges"]


def build_sales():
    wb = pyopenxlsx.Workbook()
    ws = wb.active
    ws.title = "Sales"
    headers = ["OrderID", "Month", "Region", "Product", "Quantity", "Price", "Amount"]
    ws.write_row(1, headers)

    n = 600
    for i in range(n):
        r = i + 2
        ws.set_cell_value(r, 1, 1000 + i)
        ws.set_cell_value(r, 2, MONTHS[i % 12])
        ws.set_cell_value(r, 3, REGIONS[(i // 3) % 4])
        ws.set_cell_value(r, 4, PRODUCTS[(i // 7) % 4])
        ws.set_cell_value(r, 5, random.randint(1, 50))
        ws.set_cell_value(r, 6, random.choice([250, 500, 750, 1200, 2500]))
        ws.cell(r, 7).formula = f"=E{r}*F{r}"

    total_row = n + 2
    ws.set_cell_value(total_row, 1, "Total")
    for col, letter in [(5, "E"), (6, "F"), (7, "G")]:
        ws.cell(total_row, col).formula = f"=SUM({letter}2:{letter}{n + 1})"

    ws.add_table("tblSales", f"A1:G{total_row}")
    ws.column("A").width = 10
    ws.column("B").width = 9
    ws.column("C").width = 10
    ws.column("D").width = 11
    ws.column("E").width = 10
    ws.column("F").width = 10
    ws.column("G").width = 13

    hdr = wb.add_style(font=Font(bold=True, color="FFFFFF"),
                       fill=Fill(pattern_type="solid", color="1F4E79"),
                       alignment=Alignment(horizontal="center"))
    for c in range(1, 8):
        ws.cell(1, c).style_index = hdr

    # Config sheet with a named range (TaxRate)
    cfg = wb.create_sheet("Config")
    cfg.write_row(1, ["Setting", "Value"])
    cfg.write_row(2, ["TaxRate", 0.18])
    cfg.write_row(3, ["Currency", "INR"])
    wb.defined_names.append("TaxRate", "Config!$B$2")

    # Small pivot-source sheet
    raw = wb.create_sheet("Raw")
    raw.write_row(1, ["Product", "Region", "Amount"])
    k = 2
    for p in PRODUCTS:
        for g in REGIONS:
            raw.set_cell_value(k, 1, p)
            raw.set_cell_value(k, 2, g)
            raw.set_cell_value(k, 3, random.randint(5000, 60000))
            k += 1

    # Hidden archive sheet
    arc = wb.create_sheet("Archive")
    arc.write_row(1, ["OldOrder", "Note"])
    arc.write_row(2, [9001, "migrated"])
    arc.write_row(3, [9002, "migrated"])
    arc.sheet_state = "hidden"

    path = OUT / "sales_test.xlsx"
    wb.save(str(path))
    print(f"wrote {path} sheets={wb.sheetnames}", flush=True)
    return path


def build_messy():
    wb = pyopenxlsx.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.write_row(1, ["Name", "Company", "Score", "Status"])
    rows = [
        ["Aarav", "acme corp", 78, "Active"],
        ["Diya", "Acme Corp", 92, "Active"],
        ["Kabir", "Globex", "N/A", "Pending"],
        ["Meera", "globex", 55, ""],
        ["Rohan", "Initech", 88, "Active"],
        ["Sanya", "ACME CORP", 88, "Active"],
        ["Vikram", "Hooli", "N/A", "Dropped"],
        ["Anaya", "hooli", 61, "Pending"],
        ["Ishaan", "Initech", 73, "Active"],
        ["Priya", "", 81, "Active"],
        ["Aarav", "acme corp", 78, "Active"],
        ["Riya", "Massive Dynamic", 95, "Pending"],
    ]
    for i, row in enumerate(rows):
        ws.write_row(i + 2, row)
    ws.column("A").width = 10
    ws.column("B").width = 18
    ws.column("C").width = 9
    ws.column("D").width = 10

    wb.create_sheet("Notes")  # intentionally empty (empty-read path)

    path = OUT / "messy_test.xlsx"
    wb.save(str(path))
    print(f"wrote {path} sheets={wb.sheetnames}", flush=True)
    return path


if __name__ == "__main__":
    p1 = build_sales()
    p2 = build_messy()

    # Add one cell comment via openpyxl (pyopenxlsx 1.4.2 comments crash).
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    wb = load_workbook(p1)
    wb["Sales"]["G1"].comment = Comment("Cross-check monthly totals with Raw sheet.",
                                        "AI-Test")
    wb.save(p1)
    print("comment added to sales_test.xlsx!G1", flush=True)
